"""Gerador de documentos Word (.docx) editáveis para resoluções acadêmicas manuais.

Utilizado quando o usuário dispara a resolução manualmente (/resolver, /resolver_lote),
gerando um arquivo .docx editável em vez de PDF — permitindo ao usuário revisar,
editar e aprovar antes do envio ao Moodle.
"""

import re
import sys
from pathlib import Path
from typing import Optional

from rich.console import Console

if sys.platform == "win32":
    try:
        if sys.stdout and hasattr(sys.stdout, "reconfigure"):
            sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

console = Console()


class AcademicDocxGenerator:
    """Gera documentos .docx editáveis a partir de Markdown acadêmico."""

    def __init__(self):
        try:
            import docx as _docx
            self._docx = _docx
        except ImportError:
            self._docx = None
            console.print(
                "[yellow]Aviso: python-docx não instalado. Instale com: pip install python-docx[/yellow]"
            )

    def _available(self) -> bool:
        return self._docx is not None

    def generate_docx(
        self,
        markdown_text: str,
        output_path: Path,
        course_name: str = "",
        assignment_title: str = "",
    ) -> bool:
        """Converte Markdown acadêmico em um documento .docx editável.

        Retorna True se gerado com sucesso, False em caso de erro.
        """
        if not self._available():
            console.print("[yellow]Aviso: python-docx não disponível — DOCX não gerado.[/yellow]")
            return False

        try:
            from docx import Document
            from docx.shared import Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc = Document()

            # Configuração da página (A4)
            from docx.shared import Cm
            section = doc.sections[0]
            section.page_width = Cm(21)
            section.page_height = Cm(29.7)
            section.top_margin = Cm(2.5)
            section.bottom_margin = Cm(2.5)
            section.left_margin = Cm(3)
            section.right_margin = Cm(2.5)

            # Estilo de fonte padrão
            style = doc.styles["Normal"]
            font = style.font
            font.name = "Calibri"
            font.size = Pt(12)

            # Cabeçalho discreto (matéria + título)
            if course_name or assignment_title:
                hdr_para = doc.add_paragraph()
                hdr_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if course_name:
                    run = hdr_para.add_run(course_name)
                    run.bold = True
                    run.font.size = Pt(11)
                    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                if course_name and assignment_title:
                    hdr_para.add_run("  •  ")
                if assignment_title:
                    run2 = hdr_para.add_run(assignment_title)
                    run2.italic = True
                    run2.font.size = Pt(11)
                    run2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                doc.add_paragraph()  # espaço após cabeçalho

            # Converte Markdown linha a linha
            lines = markdown_text.splitlines()
            i = 0
            while i < len(lines):
                line = lines[i]

                # Títulos
                if line.startswith("### "):
                    p = doc.add_paragraph(line[4:].strip(), style="Heading 3")
                elif line.startswith("## "):
                    p = doc.add_paragraph(line[3:].strip(), style="Heading 2")
                elif line.startswith("# "):
                    p = doc.add_paragraph(line[2:].strip(), style="Heading 1")

                # Linha horizontal
                elif line.strip() in ("---", "***", "___"):
                    p = doc.add_paragraph()
                    p.paragraph_format.border_bottom = True

                # Listas com marcador
                elif re.match(r"^[\*\-\+] ", line):
                    text = line[2:].strip()
                    p = doc.add_paragraph(style="List Bullet")
                    _add_inline_formatting(p, text)

                # Listas numeradas
                elif re.match(r"^\d+\. ", line):
                    text = re.sub(r"^\d+\. ", "", line).strip()
                    p = doc.add_paragraph(style="List Number")
                    _add_inline_formatting(p, text)

                # Linha em branco → parágrafo vazio
                elif not line.strip():
                    doc.add_paragraph()

                # Parágrafo normal
                else:
                    p = doc.add_paragraph()
                    _add_inline_formatting(p, line)

                i += 1

            output_path.parent.mkdir(parents=True, exist_ok=True)
            doc.save(str(output_path))
            console.print(f"[green]✔ DOCX gerado: {output_path.name}[/green]")
            return True

        except Exception as e:
            console.print(f"[red]Erro ao gerar DOCX: {e}[/red]")
            return False


def _add_inline_formatting(paragraph, text: str):
    """Adiciona um texto com formatação inline básica (**bold**, *italic*, `code`) a um parágrafo."""
    try:
        from docx.shared import Pt, RGBColor
    except ImportError:
        paragraph.add_run(text)
        return

    # Padrão: **bold**, *italic*, `code`, ***bold+italic***
    pattern = re.compile(r"(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*|`.*?`)")
    parts = pattern.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("***") and part.endswith("***"):
            run = paragraph.add_run(part[3:-3])
            run.bold = True
            run.italic = True
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
        else:
            paragraph.add_run(part)
